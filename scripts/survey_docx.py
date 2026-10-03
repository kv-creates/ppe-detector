"""Survey the user-edited docx structure (read-only)."""
from docx import Document

d = Document("reports/final_report.docx")
out = ["paras: %d | tables: %d" % (len(d.paragraphs), len(d.tables))]
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    tags = []
    if p.paragraph_format.page_break_before:
        tags.append("PAGE")
    if p.style.name.startswith("Heading"):
        tags.append(p.style.name)
    low = t.lower()
    if low.startswith("figure") or low.startswith("table ") or low.startswith("list of"):
        tags.append("CAPTION")
    if tags or len(t) < 110:
        out.append("%d | %s | %s" % (i, "/".join(tags), t[:85].replace("\n", " ")))
open("scripts/survey_out.txt", "w", encoding="utf-8").write("\n".join(out))
print("wrote scripts/survey_out.txt")
