"""Convert reports/final_report.md to a strictly-formatted academic DOCX.

Template rules (user-specified):
  - Times New Roman everywhere, black only (RGB 000000), no colours/highlights
  - Body 12pt, justified, 1.5 line spacing
  - Main headings 14pt bold ALL CAPS; every ## section starts on a new page
  - Sub-headings 12pt bold ALL CAPS
  - 1-inch margins, footer page numbers, field-based Table of Contents
  - Real Word tables (black grid), centred 12pt figure captions, boxed code

Run: python scripts/md_to_docx.py
Output: reports/final_report.docx
"""
from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

BLACK = RGBColor(0x00, 0x00, 0x00)
FONT = "Times New Roman"
SRC_MD = "reports/final_report.md"
OUT_DOCX = "reports/final_report.docx"
FIG_DIR = "reports"

EMOJI_RE = re.compile(
    "[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F\u2714\u2716\u274C✅❌⚠✔❐➔→←↑↓]"
)


def clean(text: str) -> str:
    text = EMOJI_RE.sub("", text)
    return re.sub(r"\s{2,}", " ", text).strip()


def set_run(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_inline(par, text: str, size=12):
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(clean(tok[2:-2])); set_run(r, size, bold=True)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(clean(tok[1:-1])); set_run(r, size, italic=True)
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(clean(tok[1:-1])); set_run(r, size)
        else:
            r = par.add_run(clean(tok)); set_run(r, size)


def base_par(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6,
             line=1.5, keep_next=False, keep_lines=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.widow_control = True
    if keep_next:
        pf.keep_with_next = True
    if keep_lines:
        pf.keep_together = True
    return p


def add_field(par, instr: str):
    r = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._element.append(f1); r._element.append(it); r._element.append(f2)
    set_run(r)


def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def add_table(doc, rows: list[list[str]]):
    if len(rows) < 2:
        return
    # drop markdown separator row (---|---|...) if present
    if len(rows) >= 2 and all(set(c.strip()) <= set("|-: ") for c in rows[1]):
        header, body = rows[0], rows[2:]
    else:
        header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.autofit = True
    for j, h in enumerate(header):
        cell = t.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean(h)); set_run(r, 12, bold=True)
        set_cell_borders(cell)
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = clean(row[j]) if j < len(row) else ""
            r = p.add_run(val); set_run(r, 12)
            set_cell_borders(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def parse_blocks(lines):
    """Yield ('h1'|'h2'|'h3'|'p'|'ul'|'ol'|'table'|'code'|'img'|'rule', payload)."""
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            buf, i = [], i + 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip()); i += 1
            i += 1
            yield ("code", "\n".join(buf))
            continue
        if s.startswith("### "):
            yield ("h3", s[4:]); i += 1; continue
        if s.startswith("## "):
            yield ("h2", s[3:]); i += 1; continue
        if s.startswith("# "):
            yield ("h1", s[2:]); i += 1; continue
        if re.match(r"^---+\s*$", s):
            yield ("rule", None); i += 1; continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            yield ("img", (m.group(1), m.group(2))); i += 1; continue
        if s.startswith("|"):
            buf = []
            while i < n and lines[i].strip().startswith("|"):
                buf.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            yield ("table", buf)
            continue
        if re.match(r"^[-*]\s+", s):
            buf = []
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                buf.append(re.sub(r"^[-*]\s+", "", lines[i].strip())); i += 1
            yield ("ul", buf)
            continue
        if re.match(r"^\d+[.)]\s+", s):
            buf = []
            while i < n and re.match(r"^\d+[.)]\s+", lines[i].strip()):
                buf.append(re.sub(r"^\d+[.)]\s+", "", lines[i].strip())); i += 1
            yield ("ol", buf)
            continue
        # paragraph: join wrapped lines
        buf = [s]; i += 1
        while i < n and lines[i].strip() and not lines[i].strip().startswith(
                ("#", "```", "|", "- ", "* ", "![", "---")) and not re.match(
                r"^\d+[.)]\s+", lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        yield ("p", " ".join(buf))


def build():
    with open(SRC_MD, encoding="utf-8") as fh:
        blocks = list(parse_blocks(fh.read().splitlines()))

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1); section.bottom_margin = Inches(1)
        section.left_margin = Inches(1); section.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = FONT; style.font.size = Pt(12)
    style.font.color.rgb = BLACK
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.5

    # footer page numbers, centred
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE")

    # ---- title page (academic project-report format) ----
    title = next((b[1] for b in blocks if b[0] == "h1"), "REPORT")
    for _ in range(3):
        doc.add_paragraph()
    p = base_par(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    r = p.add_run("PROJECT REPORT ON"); set_run(r, 14, bold=True)
    p = base_par(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    r = p.add_run(clean(title).upper()); set_run(r, 20, bold=True)
    # subtitle + author block: paragraphs until first h2
    started = False
    for kind, payload in blocks:
        if kind == "h1":
            started = True
            continue
        if kind == "h2":
            break
        if started and kind == "p":
            text = clean(payload)
            q = base_par(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
            if text.startswith("**") and text.endswith("**"):
                rr = q.add_run(text.strip("*")); set_run(rr, 14, bold=True)
            elif text.startswith("**"):
                add_inline(q, payload)
            else:
                add_inline(q, payload)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---- table of contents ----
    p = base_par(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    r = p.add_run("TABLE OF CONTENTS"); set_run(r, 14, bold=True)
    p = base_par(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')

    # ---- list of figures (from inline captions) ----
    captions = []
    for idx, (kind, payload) in enumerate(blocks):
        if kind == "img" and idx + 1 < len(blocks):
            nk, npay = blocks[idx + 1]
            if nk == "p" and npay.strip().startswith("*"):
                captions.append(clean(npay.strip().strip("*")))
    if captions:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p = base_par(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
        r = p.add_run("LIST OF FIGURES"); set_run(r, 14, bold=True)
        for cap in captions:
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.5
            p.text = ""
            add_inline(p, cap)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---- body ----
    pending_img = None
    for kind, payload in blocks:
        if kind == "h1":
            continue
        if kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run(clean(payload).upper()); set_run(r, 14, bold=True)
        elif kind == "h3":
            p = base_par(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6,
                         keep_next=True)
            r = p.add_run(clean(payload).upper()); set_run(r, 12, bold=True)
        elif kind == "p":
            # figure caption follows an image (caption text already carries
            # its own "Figure X.Y:" label, so use it verbatim)
            if pending_img is not None and payload.strip().startswith("*"):
                cap = clean(payload.strip().strip("*"))
                p = base_par(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=12,
                             keep_lines=True)
                r = p.add_run(cap); set_run(r, 12, italic=True)
                pending_img = None
            else:
                pending_img = None
                if not clean(payload):
                    continue
                p = base_par(doc)
                add_inline(p, payload)
        elif kind == "img":
            alt, rel = payload
            path = os.path.join(FIG_DIR, rel.replace("/", os.sep))
            p = base_par(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=2,
                         keep_next=True)
            if os.path.exists(path):
                p.add_run().add_picture(path, width=Inches(6.0))
            else:
                r = p.add_run(f"[missing figure: {rel}]"); set_run(r)
            pending_img = True
        elif kind == "table":
            pending_img = None
            add_table(doc, payload)
        elif kind in ("ul", "ol"):
            pending_img = None
            for item in payload:
                p = doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.5
                p.text = ""
                add_inline(p, item)
        elif kind == "code":
            pending_img = None
            for ln in payload.split("\n"):
                p = base_par(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
                r = p.add_run(ln if ln else " "); set_run(r, 10)
                pPr = p._p.get_or_add_pPr()
                bdr = OxmlElement("w:pBdr")
                for edge in ("top", "left", "bottom", "right"):
                    el = OxmlElement(f"w:{edge}")
                    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
                    el.set(qn("w:color"), "000000")
                    bdr.append(el)
                pPr.append(bdr)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
        elif kind == "rule":
            pending_img = None
            continue

    doc.save(OUT_DOCX)
    print(f"wrote {OUT_DOCX}")


if __name__ == "__main__":
    sys.exit(build())
