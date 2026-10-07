"""Repair 7.6 block order -> head, VM_A, VM_B, image, caption."""
from docx import Document

d = Document("reports/final_report.docx")
ts = [p.text for p in d.paragraphs]
i_head = ts.index("7.6 VIDEO MONITOR PAGE")
# LAST Figure 7.5 match = body caption (first match is the LoF entry)
i_cap = max(i for i, t in enumerate(ts) if t.startswith("Figure 7.5:"))
# collect [caption, img?, VM_B, VM_A] region before head (exclusive), in doc order
region = list(range(i_cap, i_head))
els = [d.paragraphs[i]._p for i in region]
head_el = d.paragraphs[i_head]._p
# classify: caption=text startswith Figure 7.5; img=empty; VM_A starts 'Upload any'; VM_B starts 'Behind'
cap = img = vma = vmb = None
for el, t in zip(els, [ts[i] for i in region]):
    if t.startswith("Figure 7.5:"):
        cap = el
    elif t == "":
        img = el
    elif t.startswith("Upload any site video"):
        vma = el
    elif t.startswith("Behind the button"):
        vmb = el
print("found:", [x is not None for x in (cap, img, vma, vmb)])
body = d.element.body
for el in els:
    body.remove(el)
anchor = head_el
for el in (vma, vmb, img, cap):
    anchor.addnext(el)
    anchor = el
d.save("reports/final_report.docx")
ts2 = [p.text for p in Document("reports/final_report.docx").paragraphs]
j = ts2.index("7.6 VIDEO MONITOR PAGE")
for t in ts2[j:j + 6]:
    print(("IMG" if t == "" else "TXT"), ":", t[:65])
