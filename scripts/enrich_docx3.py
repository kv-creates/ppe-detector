"""Additive red-zone policy note to the 5.5 TELEGRAM ALERTS docx section."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph as _P

PATH = "reports/final_report.docx"
BLACK = RGBColor(0, 0, 0)


def main():
    doc = Document(PATH)
    ts = [p.text for p in doc.paragraphs]
    i55 = ts.index("5.5 TELEGRAM ALERTS")
    # find the explanation para ("In plain words: every violation...") after it
    j = next(i for i in range(i55, len(ts))
             if ts[i].startswith("In plain words: every violation"))
    ref = doc.paragraphs[j + 1]._p  # insert after explanation, before code
    el = OxmlElement("w:p")
    ref.addprevious(el)
    p = _P(el, doc)
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Alert policy is red-zone-only: the sender checks the "
                  "detection zone first and stays silent for yellow and green "
                  "areas, so only overhead-hazard violations reach the phone. "
                  "The Video Monitor page exposes the same choice as an "
                  "\"Alert on zones\" picker (red ticked by default).")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = BLACK
    doc.save(PATH)
    print("red-zone note added")


if __name__ == "__main__":
    main()
