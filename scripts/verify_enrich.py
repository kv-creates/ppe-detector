"""Verify enrichment: original paras preserved in order + new content checks."""
from docx import Document
import zipfile

# original texts from backup
old = Document("reports/final_report_backup_before_enrich.docx") \
    if False else None
import os
bak = "reports/final_report_user_backup.docx"
new = Document("reports/final_report.docx")
old_texts = [p.text for p in Document(bak).paragraphs]
new_texts = [p.text for p in new.paragraphs]
print("backup paras:", len(old_texts), "| new paras:", len(new_texts))

# every original non-empty para must appear in order in the new file
j = 0
missing = []
for t in old_texts:
    if not t.strip():
        continue
    found = False
    while j < len(new_texts):
        if new_texts[j] == t:
            found = True
            j += 1
            break
        j += 1
    if not found:
        missing.append(t[:60])
print("original paras missing/modified:", len(missing))
for m in missing[:10]:
    print("  MISSING:", m)

added = len([t for t in new_texts if t.strip()]) - len(
    [t for t in old_texts if t.strip()])
print("net paragraphs added:", added)

# formatting of added runs: must be TNR + black
bad = 0
for p in new.paragraphs:
    if p.text in old_texts:
        continue
    for r in p.runs:
        if r.text.strip():
            if r.font.name != "Times New Roman":
                bad += 1
                print("BAD FONT:", r.text[:40], r.font.name)
            c = r.font.color.rgb
            if c is not None and str(c) != "000000":
                bad += 1
                print("BAD COLOR:", r.text[:40], c)
print("added-run violations:", bad)

heads = [p.text for p in new.paragraphs if p.text.startswith("WHAT ") or
         p.text == "5.5 STREAMLIT UI CODE"]
print("mini-heads added:", len(heads))
codes = [p.text for p in new.paragraphs
         if p.text.startswith("def derive_no_vest") or
         p.text.startswith("source = st.selectbox") or
         p.text.startswith("def win7_window")]
print("code markers present:", len(codes))
z = zipfile.ZipFile("reports/final_report.docx")
print("file valid zip, size KB:", round(os.path.getsize(
    "reports/final_report.docx") / 1024))
